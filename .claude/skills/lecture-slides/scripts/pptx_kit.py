"""강의 슬라이드용 python-pptx 도우미

    import sys; sys.path.insert(0, r"<skill>/scripts")
    from pptx_kit import Deck, C, PP_ALIGN, MSO_ANCHOR

    d = Deck(images="images", animations="animations/out")
    s = d.slide(C.BG, notes="발표자 노트")
    d.title(s, "제목")
    d.photo(s, "pizza", 0.6, 1.5, 4.6, 5.2)
    d.movie(s, "PMFBars", 6.4, 1.6, 6.33, 3.56)
    d.footer(s, d.credit("pizza"))
    d.save("output/01_제목.pptx")

좌표와 크기는 인치. 슬라이드는 13.333 × 7.5 (16:9), 여백 0.6, 제목 y 0.45, 본문 y 1.4~6.8, 하단 문구 y 6.95.
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

__all__ = ["Deck", "C", "PP_ALIGN", "MSO_ANCHOR"]


class C:
    """색상 (manim 장면과 같은 팔레트)"""
    NAVY, BG, BG2 = "1F2A44", "F7F6F2", "EEF2F8"
    BLUE, ORANGE, ORANGE_D = "2F6FDB", "F2A65A", "C8691E"
    TEXT2, GRAY, WHITE = "4A5568", "6A7489", "FFFFFF"
    CARD_DARK, LINE, ON_DARK2 = "2A3858", "C9D3E6", "9AA8C2"


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    FONT = "Malgun Gothic"

    def __init__(self, images="images", animations="animations/out", base=None):
        base = Path(base) if base else Path.cwd()
        self.img = base / images
        self.ani = base / animations
        cred = self.img / "credits.json"
        self.credits = json.loads(cred.read_text(encoding="utf-8")) if cred.exists() else {}
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)

    # 슬라이드 ---------------------------------------------------------------
    def slide(self, bg=C.BG, notes=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        print("saved", path.resolve())

    # 글자 -------------------------------------------------------------------
    def _font(self, run, size, color, bold):
        f = run.font
        f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, rgb(color), self.FONT
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):  # 한글은 동아시아 글꼴 슬롯을 따로 지정해야 적용됨
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set("typeface", self.FONT)

    def text(self, s, x, y, w, h, paras, size=24, color=C.NAVY, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.15):
        """paras: 문자열, 또는 문단 목록. 문단은 문자열이나 [(글, 크기, 색, 굵게), ...] (뒤쪽 항목 생략 가능)."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, para in enumerate([paras] if isinstance(paras, str) else paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            for r in ([(para,)] if isinstance(para, str) else para):
                t, sz, col, b = (tuple(r) + (size, color, bold)[len(r) - 1:])[:4]
                run = p.add_run()
                run.text = t
                self._font(run, sz, col, b)
        return tb

    def title(self, s, t, color=C.NAVY):
        self.text(s, 0.6, 0.45, 12.1, 0.8, t, 34, color, True)

    def footer(self, s, t, color=C.GRAY):
        self.text(s, 0.6, 6.95, 12.1, 0.35, t, 12, color)

    # 도형·사진·영상 -----------------------------------------------------------
    def box(self, s, x, y, w, h, fill, rounded=True, line=None):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            shp.adjustments[0] = 0.08
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def card(self, s, x, y, w, h, accent, head, body, head_size=22, body_size=22):
        """흰 카드 + 왼쪽 색 띠 + 제목/본문"""
        self.box(s, x, y, w, h, C.WHITE)
        self.box(s, x, y, 0.12, h, accent, rounded=False)
        self.text(s, x + 0.4, y + 0.2, w - 0.7, 0.5, head, head_size, accent, True)
        self.text(s, x + 0.4, y + 0.2 + head_size / 40, w - 0.7, h - 0.5, body, body_size, C.NAVY)

    def photo(self, s, name, x, y, w, h):
        """images/<name>.jpg 를 상자 비율에 맞게 가운데 잘라 넣기"""
        path = self.img / f"{name}.jpg"
        iw, ih = Image.open(path).size
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        img_r, box_r = iw / ih, w / h
        if img_r > box_r:
            pic.crop_left = pic.crop_right = (1 - box_r / img_r) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - img_r / box_r) / 2
        return pic

    def credit(self, name):
        """credits.json 항목. AI 생성 이미지 등은 "text" 필드에 표기 문구를 그대로 저장."""
        c = self.credits[name]
        return c.get("text") or f"사진: {c['artist']}, {c['license']}, Wikimedia Commons"

    def movie(self, s, scene, x, y, w, h, folder=None, autoplay=False):
        """<folder 또는 animations/out>/<scene>.mp4 + 포스터 <scene>.png (16:9 상자 권장: h = w × 0.5625)
        autoplay=True 이면 슬라이드가 열리자마자 재생 (기본은 클릭 재생)."""
        d = Path(folder) if folder else self.ani
        mp4 = d / f"{scene}.mp4"
        pic = s.shapes.add_movie(str(mp4), Inches(x), Inches(y), Inches(w), Inches(h),
                                 poster_frame_image=str(d / f"{scene}.png"), mime_type="video/mp4")
        if autoplay:
            self._autoplay(s, pic, _duration_ms(mp4))
        return pic

    @staticmethod
    def _autoplay(s, pic, dur_ms):
        """python-pptx가 만든 클릭 재생 타이밍을 '슬라이드와 함께 자동 재생' 타이밍으로 교체"""
        sld = s._element
        old = sld.find(qn("p:timing"))
        if old is not None:
            sld.remove(old)
        timing = parse_xml(_AUTOPLAY.format(ns=nsdecls("p"), spid=pic.shape_id, dur=dur_ms))
        ext = sld.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(timing)
        else:
            sld.append(timing)


    def table(self, s, x, y, w, h, rows, col_widths, head_fills=None, size=20, center=False):
        """rows[0]은 머리글. head_fills: 머리글 칸별 배경색 목록(기본 남색). 본문은 줄무늬."""
        t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
        head_fills = head_fills or [C.NAVY] * len(rows[0])
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                cell = t.cell(r, c)
                cell.fill.solid()
                head = r == 0
                fill = head_fills[c] if head else (C.WHITE if r % 2 else "F1F0EB")
                cell.fill.fore_color.rgb = rgb(fill)
                cell.text = v
                p = cell.text_frame.paragraphs[0]
                if center:
                    p.alignment = PP_ALIGN.CENTER
                light_head = head and fill in (C.ORANGE, C.WHITE)
                self._font(p.runs[0], size + (2 if head else 0),
                           C.BG if head and not light_head else C.NAVY, head or c == 0)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = Inches(0.15 if center else 0.25)
        for c, cw in enumerate(col_widths):
            t.columns[c].width = Inches(cw)
        return t


def _duration_ms(mp4):
    try:
        import av  # manim 설치 시 함께 설치됨
        with av.open(str(mp4)) as c:
            return int(c.duration / 1000) if c.duration else 10000
    except Exception:
        return 10000


_AUTOPLAY = """<p:timing {ns}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="{dur}" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""

"""1주차 샘플 강의자료(.pptx) 생성: manim 영상 + 비유 사진

실행: python build_pptx.py  →  output/01_확률변수와_확률분포.pptx
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "images"
ANI = HERE / "animations" / "out"
OUT = HERE / "output" / "01_확률변수와_확률분포.pptx"

FONT = "Malgun Gothic"
NAVY, BG, BG2 = "1F2A44", "F7F6F2", "EEF2F8"
BLUE, ORANGE, ORANGE_D = "2F6FDB", "F2A65A", "C8691E"
TEXT2, GRAY, WHITE = "4A5568", "6A7489", "FFFFFF"

CREDITS = json.loads((IMG / "credits.json").read_text(encoding="utf-8"))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def new_slide(bg=BG, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, size=24, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """paras: 문자열, 또는 문단 목록. 문단은 문자열이나 (글, 크기, 색, 굵게) 튜플 목록."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        runs = [(para, size, color, bold)] if isinstance(para, str) else para
        for r in runs:
            t, sz, col, b = (list(r) + [size, color, bold][len(r) - 1:])[:4]
            run = p.add_run()
            run.text = t
            f = run.font
            f.size, f.bold, f.color.rgb, f.name = Pt(sz), b, rgb(col), FONT
            rpr = run._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", FONT)
    return tb


def box(s, x, y, w, h, fill, rounded=True, line=None):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def photo(s, name, x, y, w, h):
    """사진을 상자 비율에 맞게 가운데 잘라 넣기 (object-fit: cover)"""
    path = IMG / f"{name}.jpg"
    iw, ih = Image.open(path).size
    pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    img_r, box_r = iw / ih, w / h
    if img_r > box_r:
        c = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def credit(name):
    c = CREDITS[name]
    return f"사진: {c['artist']}, {c['license']}, Wikimedia Commons"


def movie(s, scene, x, y, w, h):
    s.shapes.add_movie(str(ANI / f"{scene}.mp4"), Inches(x), Inches(y), Inches(w), Inches(h),
                       poster_frame_image=str(ANI / f"{scene}.png"), mime_type="video/mp4")


def title(s, t, color=NAVY):
    text(s, 0.6, 0.45, 12.1, 0.8, t, 34, color, True)


def footer(s, t, color=GRAY):
    text(s, 0.6, 6.95, 12.1, 0.35, t, 12, color)


# 1. 표지 ------------------------------------------------------------------
s = new_slide(NAVY, "오늘은 확률분포 단원의 첫 시간으로, 확률변수와 확률분포를 다룹니다. "
                    "오른쪽 막대 세 개는 오늘 계속 등장할 '동전 두 번 던지기'의 확률분포 모양입니다.")
for bx, bh, col in [(9.9, 1.4, BLUE), (10.8, 2.8, ORANGE), (11.7, 1.4, BLUE)]:
    box(s, bx, 6.3 - bh, 0.75, bh, col, rounded=False)
text(s, 0.8, 1.9, 8.5, 0.5, "Ⅲ. 통계 · 1. 확률분포", 20, ORANGE, True)
text(s, 0.8, 2.5, 8.5, 2.2, ["확률변수와", "확률분포"], 60, BG, True, spacing=1.05)
text(s, 0.8, 4.75, 8.5, 0.6, "우연한 결과를 숫자로 바꾸는 방법", 26, "C9D3E6")
footer(s, "확률과 통계 · 공학계열 1학년", "9AA8C2")

# 2. 학습 목표 ---------------------------------------------------------------
s = new_slide(BG, "오늘의 목표는 세 가지입니다. 확률변수가 무엇인지, 이산확률변수의 분포를 표와 그래프로 "
                  "나타내는 법, 그리고 연속확률변수에서는 확률을 왜 넓이로 구하는지입니다.")
title(s, "학습 목표")
goals = [("01", "확률변수의 뜻", "시행 결과에 실수를 대응시키는 규칙의 이해", BLUE),
         ("02", "이산확률변수", "확률분포를 표와 그래프, 확률질량함수로 표현", BLUE),
         ("03", "연속확률변수", "확률밀도함수와 넓이로 구하는 확률의 의미 파악", ORANGE_D)]
for i, (n, h, d, col) in enumerate(goals):
    x = 0.6 + i * 4.1
    box(s, x, 1.7, 3.85, 4.6, WHITE)
    box(s, x, 1.7, 3.85, 0.12, col, rounded=False)
    text(s, x + 0.4, 2.1, 3.1, 0.8, n, 40, col, True)
    text(s, x + 0.4, 3.0, 3.1, 0.7, h, 26, NAVY, True)
    text(s, x + 0.4, 3.8, 3.1, 2.0, d, 18, TEXT2, spacing=1.4)
footer(s, "교과서 Ⅲ-1 확률분포 · 확률변수와 확률분포")

# 3. 비유: 번호표 -------------------------------------------------------------
s = new_slide(BG, "확률변수를 처음 들으면 '변수'라는 말 때문에 어렵게 느낍니다. 은행 번호표를 떠올려 보세요. "
                  "손님이 누구든 번호 하나를 받습니다. 확률변수도 같습니다. 시행 결과가 무엇이든 숫자 하나를 붙여 주는 규칙입니다.")
title(s, "비유로 시작: 은행 번호표")
photo(s, "ticket", 0.6, 1.5, 4.6, 5.2)
rows = [("은행 번호표", "손님마다 번호 하나 발행", GRAY),
        ("확률변수", "시행 결과마다 실수 하나 대응", BLUE)]
for i, (h, d, col) in enumerate(rows):
    y = 1.6 + i * 1.65
    box(s, 5.7, y, 7.0, 1.4, WHITE)
    box(s, 5.7, y, 0.12, 1.4, col, rounded=False)
    text(s, 6.1, y + 0.2, 6.4, 0.5, h, 22, col, True)
    text(s, 6.1, y + 0.7, 6.4, 0.5, d, 22, NAVY)
box(s, 5.7, 5.05, 7.0, 1.4, NAVY)
text(s, 6.1, 5.05, 6.4, 1.4, [[("결과에 숫자를 붙이는 규칙", 24, BG, True)],
                              [("= ", 24, ORANGE, True), ("확률변수", 24, ORANGE, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, credit("ticket"))

# 4. 확률변수 애니메이션 --------------------------------------------------------
s = new_slide(BG, "동전을 두 번 던지면 결과는 네 가지입니다. '앞면이 몇 번 나왔나'라는 규칙으로 각 결과가 "
                  "0, 1, 2라는 숫자로 바뀝니다. (앞, 뒤)와 (뒤, 앞)이 모두 1로 모이는 점을 강조해 주세요. "
                  "그래서 X=1의 확률이 가장 큽니다. [영상: 클릭하면 재생]")
title(s, "확률변수: 결과에 숫자를 붙이는 규칙")
movie(s, "RandomVariable", 1.67, 1.35, 10.0, 5.625)
footer(s, "예: 동전 한 개를 두 번 던지는 시행 · 영상 클릭 시 재생")

# 5. 이산확률분포 ---------------------------------------------------------------
s = new_slide(BG, "앞 장의 대응을 확률로 정리하면 표와 같습니다. 네 결과가 모두 1/4씩이므로 X=1은 1/2입니다. "
                  "영상에서 막대를 눕혀 이어 붙이면 길이가 정확히 1이 됩니다. 이처럼 각 값에 확률을 대응시킨 함수가 확률질량함수입니다.")
title(s, "이산확률변수의 확률분포")
text(s, 0.6, 1.6, 5.4, 0.9, [[("값을 ", 22), ("셀 수 있는", 22, BLUE, True), (" 확률변수", 22)],
                             [("→ 이산확률변수", 22, NAVY, True)]])
tbl = s.shapes.add_table(2, 5, Inches(0.6), Inches(2.9), Inches(5.4), Inches(1.3)).table
cells = [["X", "0", "1", "2", "합계"], ["P(X=x)", "1/4", "1/2", "1/4", "1"]]
for r in range(2):
    for c in range(5):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else WHITE)
        cell.text = cells[r][c]
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        f = p.runs[0].font
        f.size, f.bold, f.name = Pt(20), r == 0, FONT
        f.color.rgb = rgb(BG if r == 0 else NAVY)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
tbl.columns[0].width = Inches(1.4)
for c in range(1, 5):
    tbl.columns[c].width = Inches(1.0)
text(s, 0.6, 4.6, 5.4, 1.2, [[("P(X = xᵢ) = pᵢ", 24, BLUE, True)], [("확률질량함수", 22, NAVY)]])
movie(s, "PMFBars", 6.4, 1.6, 6.33, 3.56)
text(s, 6.4, 5.35, 6.33, 0.5, "막대를 이어 붙이면 합계 1", 18, GRAY, align=PP_ALIGN.CENTER)
footer(s, "동전 두 번 던지기: 앞면이 나온 횟수 X의 확률분포 · 영상 클릭 시 재생")

# 6. 성질 + 피자 비유 -------------------------------------------------------------
s = new_slide(BG, "확률질량함수의 두 성질입니다. 피자 한 판을 떠올려 보세요. 조각 크기는 제각각이어도 모두 모으면 "
                  "정확히 한 판입니다. 확률도 마찬가지로 모두 더하면 1입니다. 미지수 k를 구하는 문제는 대부분 이 성질을 씁니다.")
title(s, "확률질량함수의 두 가지 성질")
photo(s, "pizza", 0.6, 1.5, 4.6, 5.2)
cards = [("성질 1", "0 ≤ pᵢ ≤ 1", "각 확률은 0 이상 1 이하", BLUE),
         ("성질 2", "p₁ + p₂ + ⋯ + pₙ = 1", "모든 확률의 합은 1", ORANGE_D)]
for i, (h, f, d, col) in enumerate(cards):
    y = 1.5 + i * 1.95
    box(s, 5.7, y, 7.0, 1.75, WHITE)
    box(s, 5.7, y, 0.12, 1.75, col, rounded=False)
    text(s, 6.1, y + 0.15, 6.4, 0.4, h, 18, col, True)
    text(s, 6.1, y + 0.55, 6.4, 0.6, f, 30, NAVY, True)
    text(s, 6.1, y + 1.2, 6.4, 0.4, d, 18, TEXT2)
box(s, 5.7, 5.5, 7.0, 1.2, NAVY)
text(s, 6.1, 5.5, 6.4, 1.2, [[("비유: ", 22, ORANGE, True), ("조각을 모두 모으면 피자 한 판", 22, BG, True)]],
     anchor=MSO_ANCHOR.MIDDLE)
footer(s, credit("pizza"))

# 7. 비유: 계단 vs 경사로 ------------------------------------------------------------
s = new_slide(BG2, "이산과 연속의 차이를 계단과 경사로로 비유합니다. 계단에서는 1층, 2층처럼 정해진 높이에만 설 수 있고 "
                   "1.5층에는 설 수 없습니다. 경사로에서는 어느 높이에서든 멈출 수 있습니다. 앞의 것이 이산, 뒤의 것이 연속입니다.")
title(s, "비유: 계단과 경사로")
for i, (name, h, d, col) in enumerate([("stairs", "이산 = 계단", "정해진 높이에만 멈춤, 중간 값 없음", BLUE),
                                       ("ramp", "연속 = 경사로", "어느 높이에서든 멈춤 가능", ORANGE_D)]):
    x = 0.6 + i * 6.2
    photo(s, name, x, 1.4, 5.9, 3.9)
    text(s, x, 5.45, 5.9, 0.55, h, 28, col, True)
    text(s, x, 6.05, 5.9, 0.5, d, 20, TEXT2)
footer(s, f"{credit('stairs')} / {credit('ramp')}")

# 8. 계단 vs 경사로 애니메이션 ----------------------------------------------------------
s = new_slide(BG, "공이 계단에서는 한 칸씩 뛰어오르고, 경사로에서는 미끄러지듯 올라갑니다. "
                  "이산확률변수는 0, 1, 2…처럼 뚝뚝 끊긴 값을, 연속확률변수는 구간 안의 모든 실수를 가집니다. [영상: 클릭하면 재생]")
title(s, "이산 vs 연속: 끊긴 값과 이어진 값")
movie(s, "StairsVsRamp", 1.67, 1.35, 10.0, 5.625)
footer(s, "영상 클릭 시 재생")

# 9. 연속확률변수 ------------------------------------------------------------------
s = new_slide(BG2, "대기 시간처럼 끝없이 잘게 나눌 수 있는 값은 셀 수 없으므로, 막대 대신 함수 그래프를 그리고 확률은 넓이로 구합니다. "
                   "높이 1/10, 폭 10이니 전체 넓이가 1입니다. 3분에서 7분 사이는 4 × 1/10 = 0.4. 구간을 한 점으로 줄이면 넓이가 0이므로 "
                   "P(X=a)=0입니다. (대기 시간 수치는 예시이므로 교과서 수치로 바꿔 쓰셔도 됩니다.)")
title(s, "연속확률변수: 확률 = 그래프 아래 넓이")
photo(s, "subway", 0.6, 1.4, 3.5, 4.4)
text(s, 0.6, 5.95, 3.5, 0.8, [[("예: 10분 간격 지하철", 18, ORANGE_D, True)], [("기다리는 시간 X", 18, NAVY, True)]])
movie(s, "DensityArea", 4.5, 1.4, 8.23, 4.63)
for i, t in enumerate(["f(x) ≥ 0", "전체 넓이 = 1", "P(X = a) = 0"]):
    b = box(s, 4.5 + i * 2.8, 6.2, 2.6, 0.6, WHITE, line="C9D3E6")
    text(s, 4.5 + i * 2.8, 6.2, 2.6, 0.6, t, 18, NAVY, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
footer(s, credit("subway") + " · 영상 클릭 시 재생")

# 10. 비교표 ---------------------------------------------------------------------
s = new_slide(BG, "두 종류를 나란히 비교합니다. 공학에서는 둘 다 자주 나옵니다. 불량품 개수는 이산, 센서가 재는 온도의 오차나 "
                  "부품 수명은 연속입니다. '몇 개냐, 얼마나냐'로 구분하게 하면 쉽게 기억합니다.")
title(s, "이산 vs 연속: 공학 현장의 예")
rows = [["구분", "이산확률변수 (계단)", "연속확률변수 (경사로)"],
        ["취하는 값", "셀 수 있는 값", "구간 안의 모든 실수"],
        ["분포 표현", "확률질량함수 P(X = x)", "확률밀도함수 f(x)"],
        ["확률 계산", "확률의 합", "그래프 아래 넓이"],
        ["공학 예시", "불량품 개수, 패킷 손실 수", "센서 측정 오차, 부품 수명"]]
tbl = s.shapes.add_table(5, 3, Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.6)).table
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        head = r == 0
        cell.fill.fore_color.rgb = rgb([NAVY, BLUE, ORANGE][c] if head else (WHITE if r % 2 else "F1F0EB"))
        cell.text = v
        p = cell.text_frame.paragraphs[0]
        f = p.runs[0].font
        f.size, f.bold, f.name = Pt(22 if head else 20), head or c == 0, FONT
        f.color.rgb = rgb(NAVY if head and c == 2 else (BG if head else NAVY))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.25)
tbl.columns[0].width = Inches(2.5)
tbl.columns[1].width = tbl.columns[2].width = Inches(4.8)
footer(s, '판단 기준: "몇 개?"면 이산, "얼마나?"면 연속')

# 11. 요약 + 퀴즈 ------------------------------------------------------------------
s = new_slide(NAVY, "정답. Q1: 이산 (1부터 6까지 셀 수 있음). Q2: X = 0, 1, 2, 3 의 확률이 각각 1/8, 3/8, 3/8, 1/8. "
                    "Q3: 폭 4 × 높이 k = 1 이므로 k = 1/4.")
title(s, "요약과 확인 퀴즈", BG)
text(s, 0.6, 1.5, 5.0, 0.5, "오늘의 핵심", 20, ORANGE, True)
text(s, 0.6, 2.1, 5.0, 2.6, ["확률변수: 결과 → 실수의 대응", "이산: 확률질량함수, 합 = 1", "연속: 확률밀도함수, 넓이 = 확률"],
     22, BG, spacing=1.9)
photo(s, "dice", 0.6, 4.6, 4.6, 2.1)
text(s, 6.2, 1.5, 6.5, 0.5, "퀴즈", 20, ORANGE, True)
for i, q in enumerate(["Q1. 주사위 한 개의 눈 X: 이산 또는 연속 판단",
                       "Q2. 동전 세 개를 던질 때 앞면의 개수 X의 확률분포표 작성",
                       "Q3. 0 ≤ X ≤ 4, f(x) = k (상수)일 때 k의 값"]):
    y = 2.1 + i * 1.55
    box(s, 6.2, y, 6.5, 1.35, "2A3858")
    text(s, 6.5, y, 5.9, 1.35, q, 20, BG, anchor=MSO_ANCHOR.MIDDLE, spacing=1.3)
footer(s, credit("dice"), "9AA8C2")

OUT.parent.mkdir(exist_ok=True)
prs.save(OUT)
print("saved", OUT)
